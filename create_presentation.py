"""Script to generate the project presentation as .pptx"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Colors
DARK_BLUE = RGBColor(0x0B, 0x1D, 0x51)
ACCENT_BLUE = RGBColor(0x1E, 0x88, 0xE5)
LIGHT_BLUE = RGBColor(0xBB, 0xDE, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MEDIUM_GRAY = RGBColor(0x66, 0x66, 0x66)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
ORANGE = RGBColor(0xEF, 0x6C, 0x00)
PURPLE = RGBColor(0x6A, 0x1B, 0x9A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height


def add_bg(slide, color=DARK_BLUE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, color, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=WHITE, bullet_char="\u2022", spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"{bullet_char} {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = spacing
    return txBox


# ── SLIDE 1: Title ──
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide, DARK_BLUE)
add_shape(slide, 0, 0, W, Inches(0.15), ACCENT_BLUE)
add_shape(slide, 0, H - Inches(0.15), W, Inches(0.15), ACCENT_BLUE)

add_text_box(slide, Inches(1), Inches(1.5), Inches(11), Inches(1.2),
             "Planejador de Viagem Personalizado", 44, WHITE, True, PP_ALIGN.CENTER)
add_text_box(slide, Inches(1), Inches(2.8), Inches(11), Inches(0.8),
             "Roteiros inteligentes com agentes de IA usando CrewAI", 24, LIGHT_BLUE, False, PP_ALIGN.CENTER)

add_shape(slide, Inches(5.5), Inches(4.0), Inches(2.3), Inches(0.05), ACCENT_BLUE)

add_text_box(slide, Inches(1), Inches(4.5), Inches(11), Inches(0.6),
             "CrewAction - Marco 2026", 20, LIGHT_BLUE, False, PP_ALIGN.CENTER)
add_text_box(slide, Inches(1), Inches(5.2), Inches(11), Inches(0.6),
             "Stack: CrewAI + Streamlit + OpenAI + SerperDev", 16, MEDIUM_GRAY, False, PP_ALIGN.CENTER)


# ── SLIDE 2: Visao Geral ──
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, 0, 0, W, Inches(1.1), DARK_BLUE)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.7),
             "Visao Geral do Projeto", 32, WHITE, True)

add_text_box(slide, Inches(0.8), Inches(1.5), Inches(11), Inches(0.8),
             "O que e?", 24, DARK_BLUE, True)
add_text_box(slide, Inches(0.8), Inches(2.2), Inches(11), Inches(1.0),
             "Uma aplicacao web que gera roteiros de viagem completos, dia a dia, "
             "utilizando 3 agentes de IA orquestrados pelo framework CrewAI em processo sequencial.",
             16, DARK_GRAY)

add_text_box(slide, Inches(0.8), Inches(3.3), Inches(11), Inches(0.8),
             "Como funciona?", 24, DARK_BLUE, True)
items = [
    "Usuario informa: destino, datas, orcamento, estilo e numero de viajantes",
    "3 agentes especializados trabalham em sequencia, passando contexto entre si",
    "Resultado: roteiro detalhado em Markdown com custos estimados e dicas praticas",
    "Interface em tempo real mostra o progresso de cada agente"
]
add_bullet_list(slide, Inches(0.8), Inches(4.0), Inches(11), Inches(3.0), items, 16, DARK_GRAY)


# ── SLIDE 3: Arquitetura ──
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, 0, 0, W, Inches(1.1), DARK_BLUE)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.7),
             "Arquitetura: CrewAI Sequential Pipeline", 32, WHITE, True)

# Flow diagram - 3 boxes
box_y = Inches(2.0)
box_h = Inches(3.5)
box_w = Inches(3.5)
gap = Inches(0.4)
start_x = Inches(1.2)

colors = [ACCENT_BLUE, GREEN, PURPLE]
titles = ["Agente 1", "Agente 2", "Agente 3"]
names = ["Pesquisador\nde Destinos", "Curador de\nExperiencias", "Redator de\nRoteiros"]
descs = [
    "Pesquisa clima,\nmoeda, transporte,\nvistos, seguranca,\nbairros e eventos",
    "Seleciona hoteis,\nrestaurantes,\natracoes e\nexperiencias alternativas",
    "Escreve roteiro\ndia a dia com\nhorarios, dicas\ne custos"
]

for i in range(3):
    x = start_x + i * (box_w + gap + Inches(0.5))
    box = add_shape(slide, x, box_y, box_w, box_h, colors[i])
    box.fill.fore_color.rgb = colors[i]

    add_text_box(slide, x + Inches(0.2), box_y + Inches(0.2), box_w - Inches(0.4), Inches(0.5),
                 titles[i], 14, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), box_y + Inches(0.6), box_w - Inches(0.4), Inches(1.0),
                 names[i], 20, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), box_y + Inches(1.8), box_w - Inches(0.4), Inches(1.5),
                 descs[i], 13, RGBColor(0xE0, 0xE0, 0xE0), False, PP_ALIGN.CENTER)

    # Arrow between boxes
    if i < 2:
        ax = x + box_w + Inches(0.05)
        add_text_box(slide, ax, box_y + Inches(1.3), Inches(0.5), Inches(0.6),
                     ">>>", 20, DARK_GRAY, True, PP_ALIGN.CENTER)

# Labels
add_text_box(slide, Inches(1.2), Inches(5.8), Inches(4), Inches(0.5),
             "Processo: Sequential  |  Contexto passado entre tarefas", 14, MEDIUM_GRAY)
add_text_box(slide, Inches(7), Inches(5.8), Inches(5), Inches(0.5),
             "Ferramenta: SerperDevTool (busca web) no Agente 1", 14, MEDIUM_GRAY)


# ── SLIDE 4: Agentes em Detalhe ──
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, 0, 0, W, Inches(1.1), DARK_BLUE)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.7),
             "Agentes e Tarefas em Detalhe", 32, WHITE, True)

y = Inches(1.4)
agents_data = [
    ("Pesquisador de Destinos", ACCENT_BLUE, [
        "Ferramenta: SerperDevTool (busca web via API SerperDev)",
        "Pesquisa: clima, moeda, cambio, transporte, vistos, seguranca",
        "Identifica bairros recomendados por estilo de viagem",
        "Mapeia eventos e festivais no periodo da viagem",
        "Output: Relatorio detalhado do destino"
    ]),
    ("Curador de Experiencias", GREEN, [
        "Recebe contexto da pesquisa do Agente 1",
        "Seleciona 2-3 opcoes de hospedagem com diarias",
        "Recomenda restaurantes variados (min. 1 por dia)",
        "Lista atracoes principais e experiencias alternativas",
        "Agrupa atividades por proximidade geografica"
    ]),
    ("Redator de Roteiros", PURPLE, [
        "Recebe contexto dos Agentes 1 e 2",
        "Formata roteiro: Dia X - [Titulo tematico]",
        "Divide cada dia: Manha / Tarde / Noite com horarios",
        "Inclui dicas praticas e custos estimados por dia",
        "Resumo final: gastos vs orcamento, checklist, frases uteis"
    ])
]

for title, color, items in agents_data:
    add_shape(slide, Inches(0.8), y, Inches(0.15), Inches(0.35), color)
    add_text_box(slide, Inches(1.1), y - Inches(0.05), Inches(4), Inches(0.5),
                 title, 18, color, True)
    for j, item in enumerate(items):
        add_text_box(slide, Inches(1.3), y + Inches(0.4) + j * Inches(0.32), Inches(11), Inches(0.35),
                     f"  {item}", 13, DARK_GRAY)
    y += Inches(2.1)


# ── SLIDE 5: Interface do Usuario ──
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, 0, 0, W, Inches(1.1), DARK_BLUE)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.7),
             "Interface do Usuario (Streamlit)", 32, WHITE, True)

# Sidebar section
add_text_box(slide, Inches(0.8), Inches(1.4), Inches(5), Inches(0.5),
             "Formulario de Entrada (Sidebar)", 20, DARK_BLUE, True)
sidebar_items = [
    "Campo de texto: destino da viagem",
    "Seletores de data: ida e volta (com validacao)",
    "Slider de orcamento: R$ 500 a R$ 500.000 (step 500)",
    "Dropdown: estilo de viagem (4 opcoes)",
    "Slider: numero de viajantes (1-10)",
    "Botao de submissao com validacao"
]
add_bullet_list(slide, Inches(0.8), Inches(2.0), Inches(5.5), Inches(3.5), sidebar_items, 15, DARK_GRAY)

# Main area section
add_text_box(slide, Inches(7), Inches(1.4), Inches(5), Inches(0.5),
             "Area Principal", 20, DARK_BLUE, True)
main_items = [
    "Tela de boas-vindas com destaques",
    "Progresso em tempo real durante execucao",
    "Status por agente: executando / concluido / erro",
    "Resultado formatado em Markdown",
    "Botao de download do roteiro",
    "Tratamento de erros com opcao de retry"
]
add_bullet_list(slide, Inches(7), Inches(2.0), Inches(5.5), Inches(3.5), main_items, 15, DARK_GRAY)

# Styles section
add_text_box(slide, Inches(0.8), Inches(5.5), Inches(11), Inches(0.5),
             "Estilos de Viagem Suportados", 18, DARK_BLUE, True)
add_text_box(slide, Inches(0.8), Inches(6.0), Inches(11), Inches(0.8),
             "Aventura (trilhas, esportes)  |  Cultura (museus, historia)  |  "
             "Gastronomia (restaurantes, mercados)  |  Relaxamento (praias, spas)",
             14, MEDIUM_GRAY, False, PP_ALIGN.LEFT)


# ── SLIDE 6: Event Listener & Concorrencia ──
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, 0, 0, W, Inches(1.1), DARK_BLUE)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.7),
             "Sistema de Eventos e Concorrencia", 32, WHITE, True)

add_text_box(slide, Inches(0.8), Inches(1.4), Inches(5.5), Inches(0.5),
             "ProgressListener (Event Listener)", 20, DARK_BLUE, True)
listener_items = [
    "Herda de BaseEventListener do CrewAI",
    "Captura eventos: CrewKickoff, AgentExecution, Task, ToolUsage",
    "ProgressState: dataclass thread-safe com Lock",
    "Atualiza status em tempo real: running / done / error",
    "Armazena resultado final e mensagens de erro"
]
add_bullet_list(slide, Inches(0.8), Inches(2.0), Inches(5.5), Inches(3.0), listener_items, 15, DARK_GRAY)

add_text_box(slide, Inches(7), Inches(1.4), Inches(5.5), Inches(0.5),
             "Modelo de Concorrencia", 20, DARK_BLUE, True)
conc_items = [
    "UI roda na thread principal do Streamlit",
    "Crew executa em thread daemon separada",
    "Comunicacao via ProgressState compartilhado",
    "Lock garante thread-safety nas escritas",
    "UI faz polling a cada 1 segundo"
]
add_bullet_list(slide, Inches(7), Inches(2.0), Inches(5.5), Inches(3.0), conc_items, 15, DARK_GRAY)

# Data flow
add_text_box(slide, Inches(0.8), Inches(5.2), Inches(11), Inches(0.5),
             "Fluxo de Dados", 20, DARK_BLUE, True)
add_text_box(slide, Inches(0.8), Inches(5.8), Inches(11), Inches(1.0),
             "Input (Streamlit) -> Validacao -> Thread(daemon) -> Crew.kickoff() -> "
             "ProgressListener captura eventos -> ProgressState atualizado -> UI renderiza resultado",
             14, MEDIUM_GRAY)


# ── SLIDE 7: Stack Tecnica ──
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape(slide, 0, 0, W, Inches(1.1), DARK_BLUE)
add_text_box(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.7),
             "Stack Tecnica e Configuracao", 32, WHITE, True)

# Tech stack
add_text_box(slide, Inches(0.8), Inches(1.4), Inches(5), Inches(0.5),
             "Dependencias", 20, DARK_BLUE, True)
deps = [
    "Python >= 3.11",
    "CrewAI[tools] >= 0.108.0",
    "Streamlit >= 1.41.0",
    "python-dotenv >= 1.0.0",
    "LLM: gpt-4o-mini (temp: 0.7)"
]
add_bullet_list(slide, Inches(0.8), Inches(2.0), Inches(5), Inches(2.5), deps, 16, DARK_GRAY)

# Config
add_text_box(slide, Inches(7), Inches(1.4), Inches(5), Inches(0.5),
             "Configuracao", 20, DARK_BLUE, True)
config = [
    "Agentes definidos em config/agents.yaml",
    "Tarefas definidas em config/tasks.yaml",
    "Variaveis de ambiente via .env",
    "OPENAI_API_KEY (obrigatorio)",
    "SERPER_API_KEY (opcional, para busca web)"
]
add_bullet_list(slide, Inches(7), Inches(2.0), Inches(5), Inches(2.5), config, 16, DARK_GRAY)

# Structure
add_text_box(slide, Inches(0.8), Inches(4.8), Inches(11), Inches(0.5),
             "Estrutura do Projeto", 20, DARK_BLUE, True)
structure = [
    "app.py — Interface Streamlit (formulario, progresso, resultado)",
    "crew.py — Fabrica de Crew (agentes, tarefas, LLM)",
    "config/agents.yaml — Configuracao dos 3 agentes (roles, goals, backstories)",
    "config/tasks.yaml — Configuracao das 3 tarefas (descricoes, outputs esperados)",
    "listeners/progress_listener.py — Event listener com ProgressState thread-safe"
]
add_bullet_list(slide, Inches(0.8), Inches(5.4), Inches(11), Inches(2.0), structure, 14, DARK_GRAY)


# ── SLIDE 8: Encerramento ──
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_shape(slide, 0, 0, W, Inches(0.15), ACCENT_BLUE)
add_shape(slide, 0, H - Inches(0.15), W, Inches(0.15), ACCENT_BLUE)

add_text_box(slide, Inches(1), Inches(1.5), Inches(11), Inches(1.0),
             "Obrigado!", 48, WHITE, True, PP_ALIGN.CENTER)

add_shape(slide, Inches(5.5), Inches(2.8), Inches(2.3), Inches(0.05), ACCENT_BLUE)

highlights = [
    "3 agentes especializados com pipeline sequencial",
    "Interface interativa com feedback em tempo real",
    "Configuracao flexivel via YAML e variaveis de ambiente",
    "Arquitetura thread-safe para execucao concorrente",
    "Roteiros completos com custos, dicas e frases uteis"
]
add_bullet_list(slide, Inches(2.5), Inches(3.3), Inches(8), Inches(3.0),
                highlights, 18, LIGHT_BLUE, "\u2714")

add_text_box(slide, Inches(1), Inches(6.3), Inches(11), Inches(0.6),
             "CrewAction - Marco 2026  |  CrewAI + Streamlit + OpenAI", 16, MEDIUM_GRAY, False, PP_ALIGN.CENTER)

# Save
output_path = "/home/user/crew-action-2026-03/Planejador_de_Viagem_CrewAI.pptx"
prs.save(output_path)
print(f"Apresentacao salva em: {output_path}")
